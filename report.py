"""
Sales Report Automator
Turns a messy sales CSV/Excel export into a clean, formatted Excel report
with summary tables and charts, in one command.

Usage:
    python report.py sample_data/sales_raw.csv
    python report.py my_export.xlsx -o June_Report.xlsx
"""
import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = '"$"#,##0.00'


# ---------- Loading & cleaning ----------

def load(path: Path) -> pd.DataFrame:
    if path.suffix.lower() in (".xlsx", ".xls"):
        return pd.read_excel(path, dtype=str)
    return pd.read_csv(path, dtype=str)


def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    df.columns = [re.sub(r"[^a-z0-9]+", "_", c.strip().lower()).strip("_") for c in df.columns]
    return df


def to_number(series: pd.Series) -> pd.Series:
    cleaned = series.astype(str).str.replace(r"[$,\s]", "", regex=True)
    return pd.to_numeric(cleaned, errors="coerce")


def parse_dates(series: pd.Series) -> pd.Series:
    parsed = pd.to_datetime(series, format="%Y-%m-%d", errors="coerce")
    for fmt in ("%m/%d/%Y", "%d-%b-%Y"):
        missing = parsed.isna()
        parsed[missing] = pd.to_datetime(series[missing], format=fmt, errors="coerce")
    missing = parsed.isna()
    if missing.any():  # last resort: let pandas guess
        parsed[missing] = pd.to_datetime(series[missing], errors="coerce")
    return parsed


def clean(df: pd.DataFrame) -> tuple[pd.DataFrame, dict, pd.DataFrame]:
    stats = {"rows_in": len(df)}
    df = normalize_columns(df)

    for col in ("region", "sales_rep", "product"):
        if col in df:
            df[col] = df[col].astype(str).str.strip().str.title()
    # Keep common abbreviations readable after title-casing
    if "product" in df:
        df["product"] = df["product"].str.replace(r"\bUsb-C\b", "USB-C", regex=True) \
                                     .str.replace(r"\bHd\b", "HD", regex=True) \
                                     .str.replace(r"\b27In\b", "27in", regex=True)

    before = len(df)
    df = df.drop_duplicates()
    stats["duplicates_removed"] = before - len(df)

    df["order_date"] = parse_dates(df["order_date"])
    df["quantity"] = to_number(df["quantity"])
    df["unit_price"] = to_number(df["unit_price"])

    bad = df[["order_date", "quantity", "unit_price"]].isna().any(axis=1)
    stats["invalid_rows_removed"] = int(bad.sum())
    rejected = df[bad].copy()
    df = df[~bad].copy()

    df["quantity"] = df["quantity"].astype(int)
    df["revenue"] = (df["quantity"] * df["unit_price"]).round(2)
    df["month"] = df["order_date"].dt.to_period("M").astype(str)
    df = df.sort_values("order_date")
    stats["rows_out"] = len(df)
    return df, stats, rejected


# ---------- Summaries ----------

def summarize(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    def group(col):
        return (df.groupby(col)
                  .agg(Orders=("order_id", "nunique"), Units=("quantity", "sum"), Revenue=("revenue", "sum"))
                  .sort_values("Revenue", ascending=False)
                  .reset_index()
                  .rename(columns={col: col.replace("_", " ").title()}))

    monthly = (df.groupby("month")
                 .agg(Orders=("order_id", "nunique"), Revenue=("revenue", "sum"))
                 .reset_index()
                 .rename(columns={"month": "Month"}))
    return {
        "By Month": monthly,
        "By Region": group("region"),
        "By Product": group("product"),
        "By Sales Rep": group("sales_rep"),
    }


# ---------- Excel output ----------

def style_sheet(ws, money_cols=()):
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        width = max(len(str(c.value or "")) for c in ws[letter]) + 3
        ws.column_dimensions[letter].width = min(width, 40)
        if ws.cell(1, col_idx).value in money_cols:
            for c in ws[letter][1:]:
                c.number_format = MONEY
    ws.freeze_panes = "A2"


def write_report(df, stats, rejected, summaries, out: Path):
    total_rev = df["revenue"].sum()
    kpis = pd.DataFrame({
        "Metric": ["Total Revenue", "Total Orders", "Units Sold", "Average Order Value",
                   "Top Product", "Top Region", "Top Sales Rep",
                   "Rows in raw file", "Duplicates removed", "Invalid rows removed", "Clean rows used"],
        "Value": [f"${total_rev:,.2f}", df["order_id"].nunique(), int(df["quantity"].sum()),
                  f"${total_rev / df['order_id'].nunique():,.2f}",
                  summaries["By Product"].iloc[0, 0], summaries["By Region"].iloc[0, 0],
                  summaries["By Sales Rep"].iloc[0, 0],
                  stats["rows_in"], stats["duplicates_removed"], stats["invalid_rows_removed"],
                  stats["rows_out"]],
    })

    clean_out = df.drop(columns=["month"]).copy()
    clean_out["order_date"] = clean_out["order_date"].dt.date
    clean_out.columns = [c.replace("_", " ").title() for c in clean_out.columns]

    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        kpis.to_excel(xw, sheet_name="Summary", index=False)
        for name, table in summaries.items():
            table.to_excel(xw, sheet_name=name, index=False)
        clean_out.to_excel(xw, sheet_name="Clean Data", index=False)
        if len(rejected):
            rejected.to_excel(xw, sheet_name="Rejected Rows", index=False)

    wb = load_workbook(out)
    for ws in wb.worksheets:
        style_sheet(ws, money_cols=("Revenue", "Unit Price"))

    # Charts
    ws = wb["By Month"]
    line = LineChart()
    line.title, line.y_axis.title, line.height, line.width = "Monthly Revenue", "Revenue ($)", 8, 16
    line.add_data(Reference(ws, min_col=3, min_row=1, max_row=ws.max_row), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    ws.add_chart(line, "F2")

    for sheet in ("By Product", "By Region", "By Sales Rep"):
        ws = wb[sheet]
        bar = BarChart()
        bar.type, bar.title, bar.height, bar.width = "bar", f"Revenue {sheet.lower()}", 8, 16
        bar.add_data(Reference(ws, min_col=4, min_row=1, max_row=ws.max_row), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        bar.legend = None
        ws.add_chart(bar, "G2")

    wb.save(out)


def main():
    ap = argparse.ArgumentParser(description="Turn a messy sales export into a clean Excel report.")
    ap.add_argument("input", type=Path, help="CSV or Excel file")
    ap.add_argument("-o", "--output", type=Path, default=Path("sales_report.xlsx"))
    args = ap.parse_args()

    if not args.input.exists():
        sys.exit(f"File not found: {args.input}")

    df, stats, rejected = clean(load(args.input))
    if df.empty:
        sys.exit("No valid rows found after cleaning.")
    write_report(df, stats, rejected, summarize(df), args.output)

    print(f"Report saved to {args.output}")
    print(f"  {stats['rows_in']} rows in -> {stats['rows_out']} clean rows "
          f"({stats['duplicates_removed']} duplicates, {stats['invalid_rows_removed']} invalid removed)")
    print(f"  Total revenue: ${df['revenue'].sum():,.2f}")


if __name__ == "__main__":
    main()
